import docx
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DOC_PATH = "Carpeta_Turnify_Final_Oficial_IPET1308_Definitiva.docx"

def set_cell_margins_and_shading(cell, fill_hex=None):
    tcPr = cell._tc.get_or_add_tcPr()
    # Margins (padding)
    tcMar = parse_xml(
        '<w:tcMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '  <w:top w:w="80" w:type="dxa"/>'
        '  <w:bottom w:w="80" w:type="dxa"/>'
        '  <w:left w:w="120" w:type="dxa"/>'
        '  <w:right w:w="120" w:type="dxa"/>'
        '</w:tcMar>'
    )
    tcPr.append(tcMar)
    if fill_hex:
        shd = parse_xml(
            f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            f'w:val="clear" w:color="auto" w:fill="{fill_hex}"/>'
        )
        tcPr.append(shd)

def update_gantt_chart():
    doc = docx.Document(DOC_PATH)

    # 1. Actualizar párrafo introductorio (P143)
    p_intro = doc.paragraphs[143]
    p_intro.text = (
        "El cronograma de trabajo planifica y registra las tareas del proyecto a lo largo del ciclo lectivo 2026, "
        "abarcando desde abril hasta diciembre. Se contemplan todas las fases del ciclo de vida del software, "
        "desde la concepción y análisis inicial hasta el despliegue continuo, auditorías de calidad y la defensa final:"
    )
    if p_intro.runs:
        p_intro.runs[0].font.name = "Calibri"
        p_intro.runs[0].font.size = Pt(11)

    # 2. Actualizar nota al pie de la figura (P145)
    p_nota = doc.paragraphs[145]
    p_nota.text = "Nota. Las marcas (X) indican los periodos planificados y ejecutados para cada fase del proyecto durante el ciclo lectivo 2026 (abril a diciembre)."
    if p_nota.runs:
        p_nota.runs[0].font.name = "Calibri"
        p_nota.runs[0].font.size = Pt(9)
        p_nota.runs[0].font.italic = True

    # 3. Datos del nuevo Gantt (Abril a Diciembre = 9 meses)
    headers = ["Actividad / Fase", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
    
    rows_data = [
        ("1. Análisis de requerimientos y estudio de factibilidad",      ["X", "X", "",  "",  "",  "",  "",  "",  ""]),
        ("2. Diseño de arquitectura, casos de uso y modelo de datos",  ["",  "X", "X", "",  "",  "",  "",  "",  ""]),
        ("3. Configuración del entorno base y persistencia Firestore",  ["",  "",  "X", "X", "",  "",  "",  "",  ""]),
        ("4. Desarrollo del panel de cliente (reserva y autogestión)",  ["",  "",  "",  "X", "X", "",  "",  "",  ""]),
        ("5. Desarrollo del panel de administración y agenda SaaS",     ["",  "",  "",  "",  "X", "X", "",  "",  ""]),
        ("6. Módulo de control de caja, cobros y cierres de turno",     ["",  "",  "",  "",  "",  "X", "X", "",  ""]),
        ("7. Pruebas unitarias, control de calidad y corrección de bugs", ["", "", "",  "",  "",  "X", "X", "X", ""]),
        ("8. Despliegue en producción serverless (Vercel) y ajustes UI",["",  "",  "",  "",  "",  "",  "X", "X", ""]),
        ("9. Documentación técnica, carpeta de campo y presentación",   ["X", "X", "X", "X", "X", "X", "X", "X", "X"]),
    ]

    col_widths = [Inches(2.7)] + [Inches(0.465)] * 9

    # 4. Crear nueva tabla
    new_table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    new_table.style = "Table Grid"
    new_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Encabezado
    hdr_cells = new_table.rows[0].cells
    for c_idx, title in enumerate(headers):
        hdr_cells[c_idx].text = title
        set_cell_margins_and_shading(hdr_cells[c_idx], "E2E2E2")
        hdr_cells[c_idx].width = col_widths[c_idx]
        p = hdr_cells[c_idx].paragraphs[0]
        if c_idx > 0:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.bold = True

    # Filas de datos
    for r_idx, (act, months) in enumerate(rows_data):
        row_cells = new_table.rows[r_idx + 1].cells
        # Columna de actividad
        row_cells[0].text = act
        row_cells[0].width = col_widths[0]
        set_cell_margins_and_shading(row_cells[0], None)
        p_act = row_cells[0].paragraphs[0]
        p_act.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p_act.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            run.font.bold = False

        # Columnas de meses (Abr a Dic)
        for m_idx, mark in enumerate(months):
            c = row_cells[m_idx + 1]
            c.text = mark
            c.width = col_widths[m_idx + 1]
            shading = "D9E1F2" if mark == "X" else None
            set_cell_margins_and_shading(c, shading)
            p_m = c.paragraphs[0]
            p_m.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p_m.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(8.5)
                run.font.bold = True if mark == "X" else False
                if mark == "X":
                    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    # 5. Reemplazar la tabla 10 original en el árbol XML
    old_table = doc.tables[10]
    old_element = old_table._element
    new_element = new_table._element
    old_element.getparent().replace(old_element, new_element)

    output_path = "Carpeta_Turnify_Final_Oficial_IPET1308_Gantt_Actualizado.docx"
    doc.save(output_path)
    print(f"Documento guardado con éxito como '{output_path}'.")
    
    # Intentar sobreescribir el definitivo si no está bloqueado
    try:
        import shutil
        shutil.copyfile(output_path, DOC_PATH)
        print(f"Sobrescrito exitosamente en '{DOC_PATH}'.")
    except Exception as e:
        print(f"Nota: '{DOC_PATH}' está abierto en Microsoft Word. El nuevo documento quedó generado como '{output_path}'. Cierre Word si desea sobrescribir el archivo original.")

if __name__ == "__main__":
    update_gantt_chart()
