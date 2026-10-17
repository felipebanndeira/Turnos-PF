import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOC_IN = "Carpeta_Turnify_Final_Oficial_IPET1308_Gantt_Actualizado.docx"
DOC_OUT = "Carpeta_Turnify_Final_Oficial_IPET1308_Definitiva_Con_Pantallas.docx"

pantallas_data = [
    {
        "num": "Figura 4",
        "titulo": "Pantalla P-01: Portal de Bienvenida y Acceso al Sistema (Landing)",
        "desc": (
            "Pantalla P-01 (Landing / Acceso): Portal inicial con diseño responsivo bajo estética Dark Theme "
            "con acentos azul eléctrico. Dispone de acceso directo para que los clientes soliciten turnos en sus comercios "
            "habituales sin requerir registro previo, y un enlace de autenticación segura exclusivo para el administrador."
        ),
        "img": "capturas_pantallas/p01_landing.png",
        "width": Inches(5.8),
        "nota": "Nota. Interfaz pública principal de bienvenida y acceso al sistema."
    },
    {
        "num": "Figura 5",
        "titulo": "Pantalla P-02: Flujo de Reserva Mobile-First para Clientes",
        "desc": (
            "Pantalla P-02 (Solicitud de Turno - Reserva Cliente): Proceso de reserva mobile-first guiado en 3 pasos con "
            "barra de progreso e indicadores visuales. Incluye selector interactivo de servicios con tarifas y duraciones, "
            "carrusel horizontal de fechas con validación de horarios en tiempo real contra Firestore, y barra flotante de conversión."
        ),
        "img": "capturas_pantallas/p02_solicitar_turno.png",
        "width": Inches(3.2),  # mobile view
        "nota": "Nota. Vista interactiva en dispositivo móvil optimizada para reserva ágil en 3 pasos."
    },
    {
        "num": "Figura 6",
        "titulo": "Pantalla P-03: Consulta y Autogestión de Citas (Mis Turnos)",
        "desc": (
            "Pantalla P-03 (Mis Turnos / Estado de Cita): Portal de autoservicio para el cliente que permite consultar citas "
            "vigentes o históricas mediante su número telefónico. Brinda la opción de verificar el estado (pendiente, confirmado, cobrado) "
            "y autogestionar la cancelación anticipada respetando las políticas horarias configuradas por el comercio."
        ),
        "img": "capturas_pantallas/p03_mis_turnos.png",
        "width": Inches(3.2),  # mobile view
        "nota": "Nota. Módulo de autoservicio y seguimiento de reservas para el cliente."
    },
    {
        "num": "Figura 7",
        "titulo": "Pantalla P-04: Panel de Control y Agenda Operativa de Turnos",
        "desc": (
            "Pantalla P-04 (Agenda de Administración): Tablero operativo diario con vista cronológica tipo timeline y tabla "
            "filtrable por estado (Pendiente, Confirmado, Cancelado, Cobrado). Permite al comerciante confirmar turnos, reprogramar "
            "fecha y horario con verificación automática de solapamientos, enviar recordatorios por WhatsApp y registrar pagos."
        ),
        "img": "capturas_pantallas/p04_agenda.png",
        "width": Inches(5.8),
        "nota": "Nota. Panel de gestión operativa de turnos con métricas de estado y acciones rápidas."
    },
    {
        "num": "Figura 8",
        "titulo": "Pantalla P-05: Módulo Financiero, Registro de Cobros y Arqueo de Caja",
        "desc": (
            "Pantalla P-05 (Caja y Arqueo Diario): Panel de conciliación financiera y arqueo de jornada. Muestra tarjetas métricas "
            "estilo SaaS con el total recaudado, discriminación en efectivo y transferencias electrónicas, control de turnos pendientes "
            "de cobro y auditoría histórica con opción de snapshot y reinicio diario."
        ),
        "img": "capturas_pantallas/p05_caja.png",
        "width": Inches(5.8),
        "nota": "Nota. Panel de control de caja diaria y conciliación de métodos de pago."
    },
    {
        "num": "Figura 9",
        "titulo": "Pantalla P-06: Configuración del Negocio, Catálogo de Servicios y Horarios",
        "desc": (
            "Pantalla P-06 (Configuración del Negocio y Horarios): Interfaz modular avanzada que permite al comerciante "
            "definir el nombre, slug público de reservas, WhatsApp de contacto, catálogo dinámico de servicios (precios y duraciones), "
            "franjas de atención semanal, descansos periódicos y bloqueo de feriados con selector de calendario interactivo."
        ),
        "img": "capturas_pantallas/p06_configuracion.png",
        "width": Inches(5.8),
        "nota": "Nota. Panel de configuración comercial integral sin requerir soporte técnico."
    }
]

def add_screenshots_to_word():
    doc = docx.Document(DOC_IN)

    # Localizar sección 6.2 (después del índice de contenidos)
    idx_start = None
    for i, p in enumerate(doc.paragraphs):
        if i > 50 and "6.2. Frontend / Prototipo" in p.text:
            idx_start = i
            break

    if idx_start is None:
        raise ValueError("No se encontró la sección 6.2 en el documento.")

    print(f"Sección 6.2 encontrada en párrafo {idx_start}")

    # Párrafo introductorio (6.2)
    p_intro = doc.paragraphs[idx_start + 1]
    p_intro.text = (
        "El prototipo visual del sistema fue diseñado e implementado siguiendo rigurosos principios de experiencia de usuario (UX/UI), "
        "priorizando la fluidez y simplicidad de reserva para los clientes en dispositivos móviles (mobile-first) "
        "y la máxima claridad operativa para el comerciante desde el panel de control. "
        "A continuación se documentan las pantallas principales que componen la plataforma Turnify, acompañadas de su especificación técnica y funcional:"
    )
    p_intro.runs[0].font.name = "Calibri"
    p_intro.runs[0].font.size = Pt(11)

    # Identificar el párrafo antes de 6.3 para reemplazar las viñetas P-01 a P-05
    idx_end = None
    for i in range(idx_start + 2, min(idx_start + 15, len(doc.paragraphs))):
        if "6.3. Tabla de trazabilidad" in doc.paragraphs[i].text:
            idx_end = i
            break

    if idx_end is None:
        raise ValueError("No se encontró el inicio de la sección 6.3.")

    print(f"Borrando viñetas antiguas entre {idx_start + 2} y {idx_end}...")
    
    # Eliminar párrafos viejos de P-01 a P-05 (en orden inverso)
    for i in range(idx_end - 1, idx_start + 1, -1):
        p_to_remove = doc.paragraphs[i]._element
        p_to_remove.getparent().remove(p_to_remove)

    # Referencia al elemento previo (la intro de 6.2)
    current_element = doc.paragraphs[idx_start + 1]._element

    # Insertar cada pantalla con su título APA, descripción, imagen y nota
    for pdata in pantallas_data:
        # 1. Párrafo descriptivo
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(8)
        p_desc.paragraph_format.space_after = Pt(4)
        run_desc = p_desc.add_run(pdata["desc"])
        run_desc.font.name = "Calibri"
        run_desc.font.size = Pt(11)
        current_element.addnext(p_desc._element)
        current_element = p_desc._element

        # 2. Título de Figura (APA 7)
        p_fig = doc.add_paragraph()
        p_fig.paragraph_format.space_before = Pt(6)
        p_fig.paragraph_format.space_after = Pt(2)
        r1 = p_fig.add_run(pdata["num"])
        r1.font.name = "Calibri"
        r1.font.bold = True
        r1.font.size = Pt(10)
        
        r2 = p_fig.add_run(f"\n{pdata['titulo']}")
        r2.font.name = "Calibri"
        r2.font.bold = True
        r2.font.size = Pt(10)
        current_element.addnext(p_fig._element)
        current_element = p_fig._element

        # 3. Imagen centrada
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(2)
        r_img = p_img.add_run()
        r_img.add_picture(pdata["img"], width=pdata["width"])
        current_element.addnext(p_img._element)
        current_element = p_img._element

        # 4. Nota al pie de figura (APA 7)
        p_nota = doc.add_paragraph()
        p_nota.paragraph_format.space_before = Pt(2)
        p_nota.paragraph_format.space_after = Pt(12)
        r_nota = p_nota.add_run(pdata["nota"])
        r_nota.font.name = "Calibri"
        r_nota.font.size = Pt(9)
        r_nota.font.italic = True
        current_element.addnext(p_nota._element)
        current_element = p_nota._element

    doc.save(DOC_OUT)
    print(f"Documento generado exitosamente como '{DOC_OUT}'.")

    # Intentar sobrescribir Carpeta_Turnify_Final_Oficial_IPET1308_Definitiva.docx si no está bloqueada
    try:
        import shutil
        shutil.copyfile(DOC_OUT, "Carpeta_Turnify_Final_Oficial_IPET1308_Definitiva.docx")
        print("También se actualizó 'Carpeta_Turnify_Final_Oficial_IPET1308_Definitiva.docx'.")
    except Exception as e:
        print("Nota: 'Carpeta_Turnify_Final_Oficial_IPET1308_Definitiva.docx' está abierta en Word. Cierre Word para sobrescribirla automáticamente.")

if __name__ == "__main__":
    add_screenshots_to_word()
